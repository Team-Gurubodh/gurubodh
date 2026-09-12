"""Offline command execution against a non-editable install, never checkout code.

Run with python -I from outside the checkout. Only SDK calls and model loading
are substituted; composition, conversion, pipelines and artifact I/O stay real.
"""

from contextlib import ExitStack, chdir, redirect_stderr, redirect_stdout
import hashlib
from importlib.metadata import distribution
from io import StringIO
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document
import gurubodh
from gurubodh.cli import COMPOSED_COMMANDS, main
from gurubodh.legacy.converter import convert_texts
from gurubodh.resource_discovery import bundled_resource_path


class FixtureTokenizer:
    def encode(self, text, add_special_tokens=False):
        return list(range(len(text)))


class FixtureModel:
    tokenizer = FixtureTokenizer()

    def encode(self, texts, **kwargs):
        return [[1.0, float(index % 2)] for index, _ in enumerate(texts)]


def generate_content(*, contents, model, config):
    assert model == "gemini-3.6-flash"
    assert config.response_mime_type == "application/json"
    assert contents.startswith("<source-text>\n") and contents.endswith("\n</source-text>")
    source = contents.removeprefix("<source-text>\n").removesuffix("\n</source-text>")
    return SimpleNamespace(
        text=json.dumps({"corrected_text": source, "edits": []}),
        usage_metadata=None,
    )


def invoke(argv, *, error=None):
    stdout, stderr = StringIO(), StringIO()
    with redirect_stdout(stdout), redirect_stderr(stderr):
        try:
            result = main([str(value) for value in argv])
        except SystemExit as exc:
            assert error is not None and exc.code == 2, (argv, exc.code, stderr.getvalue())
        else:
            assert error is None and result is None, (argv, result, stderr.getvalue())
    if error:
        assert error in stderr.getvalue(), (argv, stderr.getvalue())
        assert "Traceback" not in stderr.getvalue()
    else:
        assert not stderr.getvalue() or argv[0] == "compare-tokenizers", stderr.getvalue()
    return stdout.getvalue()


def verify_install(fixtures):
    module = Path(gurubodh.__file__).resolve()
    installed = distribution("gurubodh_cli")
    direct_url = json.loads(installed.read_text("direct_url.json") or "{}")
    assert not direct_url.get("dir_info", {}).get("editable", False)
    assert not (module.parents[1] / "pyproject.toml").exists(), module
    assert module == Path(installed.locate_file("gurubodh/__init__.py")).resolve(), module
    assert sys.flags.isolated, "Use python -I to prevent checkout/PYTHONPATH imports"
    assert Path.cwd() != module.parent
    assert os.getuid() != 0, "Run as the image's normal non-root user"

    golden = json.loads((fixtures / "aps_prakash_golden.json").read_text())
    wrapper = bundled_resource_path("scripts/legacy_font_convert.js")
    vendor = bundled_resource_path(golden["mapping"]["vendor_file"])
    assert vendor.parent == wrapper.parent / "vendor"
    assert hashlib.sha256(vendor.read_bytes()).hexdigest() == golden["mapping"]["sha256"]
    assert convert_texts([case["legacy_input"] for case in golden["cases"]], "aps", wrapper) == [
        case["expected_unicode"] for case in golden["cases"]
    ]
    print(f"Installed package: {module}; APS golden conversions passed.")


def write_subject(root, fixtures, encoding, language):
    manifest = json.loads((fixtures / "job-components/subjects/aps-hindi.json").read_text())
    edition = manifest["editions"]["hi-IN"]
    edition["source_document"]["font_encoding"] = encoding
    manifest["editions"] = {language: edition}
    manifest_dir = root / "project/jobs/subjects" / manifest["manifest_id"]
    manifest_dir.mkdir(parents=True)
    (manifest_dir / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    source = root / "source" / edition["source_document"]["relative_path"]
    source.parent.mkdir(parents=True)
    document = Document()
    document.styles["Normal"].font.name = "Mangal"
    document.add_paragraph("प्रबोधन १")
    run = document.add_paragraph().add_run("efkeâ&" if encoding == "aps" else "ज्ञान")
    run.font.name = "APS-DV-Prakash" if encoding == "aps" else "Mangal"
    document.add_paragraph("पहला वाक्य। दूसरा वाक्य। तीसरा वाक्य।")
    document.save(source)
    return manifest, source


def exercise_commands(root, fixtures, encoding, language):
    manifest, source = write_subject(root, fixtures, encoding, language)
    source_bytes = source.read_bytes()
    selectors = ["--project-root", root / "project", "--subject", manifest["manifest_id"],
                 "--language", language, "--environment", "development", "--storage-profile", "local"]
    environment = {"GURUBODH_SOURCE_LIBRARY_ROOT": str(root / "source"),
                   "GURUBODH_CMS_LIBRARY_ROOT": str(root / "artifacts"),
                   "GEMINI_API_KEY": "offline-fixture-key"}
    with patch.dict(os.environ, environment):
        for command in COMPOSED_COMMANDS:
            resolved = json.loads(invoke(["config", "resolve", "--command", command, *selectors]))
            assert resolved["source"]["backend"] == resolved["destination"]["backend"] == "local"

        invoke(["prep-subject", *selectors])
        subject = root / "artifacts" / manifest["artifact_root"] / language
        canonical = list((subject / "chapters/text_and_metadata").glob("*.txt"))
        assert len(canonical) == 1, list(subject.rglob("*"))
        text = canonical[0].read_text(encoding="utf-8")
        assert ("र्कि" if encoding == "aps" else "ज्ञान") in text, text
        assert "पहला वाक्य। दूसरा वाक्य। तीसरा वाक्य।" in text
        checkpoint = json.loads((subject / "run_state/prep-subject/job-state.json").read_text())
        assert checkpoint["state"] == checkpoint["publication"]["state"] == "succeeded"
        candidate = json.loads((subject / "chapters/chapter_content_manifest.json").read_text())
        assert len(candidate["chapters"]) == 1

        invoke(["generate-chunks", *selectors])
        chunks = list((subject / "chapters").rglob("*.chunks.json"))
        assert len(chunks) == 1, list(subject.rglob("*"))
        chunk_payload = json.loads(chunks[0].read_text())
        assert chunk_payload["chunks"]
        assert "".join("".join(c["text"].split()) for c in chunk_payload["chunks"]) == "".join(text.split())

        invoke(["generate-docx", *selectors])
        exports = list((subject / "chapters/msword").glob("*.docx"))
        assert len(exports) == 1, list(subject.rglob("*"))
        paragraphs = [p.text for p in Document(exports[0]).paragraphs]
        assert paragraphs[0] == "aps-example: prabodhan 001", paragraphs
        assert "\n\n".join(paragraphs[1:]) + "\n" == text

        assembly = root / "assembled.docx"
        invoke(["lab", "assemble-docx", exports[0].parent, assembly])
        assert [p.text for p in Document(assembly).paragraphs if p.text] == [p for p in paragraphs if p]
        invoke(["lab", "append-docx", exports[0], assembly, "--no-page-break"])
        assert [p.text for p in Document(assembly).paragraphs if p.text] == [p for p in paragraphs if p] * 2

        invoke(["lab", "proofread", "--source", source, "--locale", language,
                "--lab-root", root / "lab", "--project-root", root / "project"])
        reports = list((root / "lab").glob("**/succeeded/*/run_manifest.json"))
        assert len(reports) == 1, list((root / "lab").rglob("*"))
        report = json.loads(reports[0].read_text())
        assert report["run_identity"]["status"] == "succeeded"
        assert report["publication"]["canonical"] is False

        comparison = json.loads(invoke(["compare-tokenizers", "--source-file", canonical[0],
                                       "--local-files-only", "--format", "json"]))
        assert comparison["comparison_count"] == 1
        assert comparison["files"][0]["bge"]["token_count"] == len("".join(text.split()))
        invoke(["compare-tokenizers", "--source-file", canonical[0], "--include-sarvam"],
               error="--approve-external-api")
        # The real writer must refuse accidental replacement.
        invoke(["generate-docx", *selectors], error="overwrite")
        invoke(["generate-docx", *selectors, "--overwrite"])
        assert source.read_bytes() == source_bytes
        assert canonical[0].read_text(encoding="utf-8") == text
    print(f"Installed {encoding}/{language}: prep, chunks, DOCX, all lab commands, tokenizer and negative paths passed.")


def run():
    fixtures = Path(__file__).resolve().parent / "fixtures"
    verify_install(fixtures)
    # Import real dependencies before replacing only network/model entry points.
    import boto3
    from google.genai.models import Models
    from sentence_transformers import SentenceTransformer
    from transformers import AutoTokenizer
    assert callable(boto3.client)

    with tempfile.TemporaryDirectory(prefix="gurubodh-installed-runtime-") as directory, ExitStack() as stack:
        root = Path(directory)
        stack.enter_context(chdir(root))
        stack.enter_context(patch("socket.create_connection", side_effect=AssertionError("unexpected network")))
        stack.enter_context(patch("socket.socket.connect", side_effect=AssertionError("unexpected network")))
        gemini = stack.enter_context(patch.object(Models, "generate_content", side_effect=generate_content))
        embeddings = stack.enter_context(patch("sentence_transformers.SentenceTransformer", return_value=FixtureModel()))
        tokenizer = stack.enter_context(patch.object(AutoTokenizer, "from_pretrained", return_value=FixtureTokenizer()))
        for encoding, language in (("aps", "hi-IN"), ("unicode", "mr-IN")):
            exercise_commands(root / encoding, fixtures, encoding, language)
        assert gemini.call_count == 4, gemini.call_count
        assert embeddings.call_count == 2, embeddings.call_count
        assert tokenizer.call_count == 2, tokenizer.call_count
        # Exercise the installed console-script entry point as a subprocess too.
        for argv in (["--help"], ["prep-subject", "--help"], ["lab", "--help"]):
            subprocess.run(["gurubodh", *argv], check=True, capture_output=True, text=True)
    print("Offline installed-runtime gate passed; no live provider/R2 calls or model weights used.")


if __name__ == "__main__":
    run()
