"""Issue #307: shared approvals, strict failures, and both command boundaries."""

from contextlib import chdir
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from docx import Document

from check_installed_runtime import generate_content, invoke, write_subject
from gurubodh.errors import ConfigurationError
from gurubodh.legacy.font_detection import (
    UnsupportedSourceFontError,
    detect_converter_for_font,
    load_approved_unicode_font_families,
    validate_supported_source_fonts,
)
from gurubodh.resource_discovery import BundledResourceError, bundled_resource_path
from gurubodh.schema_validation import _validator, validate_component


CLI_ROOT = Path(__file__).parents[1]
POLICY_PATH = "config/policies/source-fonts.json"


class SourceFontPolicyTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.policy = json.loads(bundled_resource_path(POLICY_PATH).read_text())
        self.path = self.root / "source-fonts.json"
        self.write_policy()

    def write_policy(self):
        self.path.write_text(json.dumps(self.policy), encoding="utf-8")

    def policy_resource(self):
        return patch("gurubodh.legacy.font_detection.bundled_resource_path", return_value=self.path)

    def source(self, family):
        path = self.root / "source.docx"
        if family is None:
            with zipfile.ZipFile(path, "w") as package:
                package.writestr("word/document.xml", (
                    '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                    '<w:body><w:p><w:r><w:t>ज्ञान</w:t></w:r></w:p></w:body></w:document>'
                ))
            return path
        document = Document()
        document.add_paragraph("ज्ञान").runs[0].font.name = family
        document.save(path)
        return path

    def test_preserves_all_original_unicode_approvals(self):
        original = {
            "aparajita", "aptos", "arial", "arial unicode ms", "calibri", "cambria",
            "hind", "kohinoor devanagari", "kokila", "lohit devanagari", "mangal",
            "mukta", "nirmala ui", "noto sans devanagari", "noto serif devanagari",
            "sanskrit 2003", "shobhika", "times new roman", "tiro devanagari hindi", "utsaah",
        }
        self.assertTrue(original.issubset(load_approved_unicode_font_families()))
        for family in original:
            with self.subTest(family=family):
                validate_supported_source_fonts(self.source(family))

    def test_matches_complete_names_after_case_and_whitespace_normalization(self):
        self.policy["approved_unicode_font_families"].append("  New\tUnicode  FAMILY ")
        self.write_policy()
        with self.policy_resource():
            validate_supported_source_fonts(self.source("new unicode family"))
            validate_supported_source_fonts(self.source("  NOTO   Sans Devanagari "))
            for family in ("New Unicode", "New Unicode Family Extra", "Other New Unicode Family",
                           "Mangal Extra", "NotMangal", "Arial Unicode", "Noto Sans"):
                with self.subTest(family=family), self.assertRaises(UnsupportedSourceFontError):
                    validate_supported_source_fonts(self.source(family))

    def test_rejects_every_known_legacy_classification_in_unicode_policy(self):
        for family in ("APS-DV-Prakash", "APS DV Test", "APS_DV_Test", "Priyanka", "Prakash",
                       "ShreeLipi", "Shree-Lipi", "Shree Lipi", "SriLipi", "Sri-Lipi",
                       "Sri Lipi", "SHREE-DEV7-0708", "ShreeDev", "  APS   DV Test "):
            with self.subTest(family=family), self.policy_resource():
                self.policy["approved_unicode_font_families"] = [family]
                self.write_policy()
                with self.assertRaisesRegex(ConfigurationError, "conflicts.*legacy-font classification"):
                    load_approved_unicode_font_families()

    def test_legacy_converter_selection_is_unchanged(self):
        for family in ("APS-DV-Test", "APS DV Test", "APS_DV_Test", "Priyanka", "Prakash"):
            with self.subTest(family=family):
                self.assertEqual(detect_converter_for_font(family), "aps")
                validate_supported_source_fonts(self.source(family))
        for family in (None, "Mangal", "ShreeLipi", "SHREE-DEV7-0708", "ShreeLipi APS-DV"):
            with self.subTest(family=family):
                self.assertIsNone(detect_converter_for_font(family))

    def test_strict_schema_rejects_invalid_policy_shapes(self):
        invalid = [None, [], {}, {"schema_version": "1.0.0"},
                   {"approved_unicode_font_families": ["Mangal"]},
                   {**self.policy, "schema_version": "2.0.0"},
                   {**self.policy, "subject": "override"}]
        for families in (None, "Mangal", [], [None], [42], [{}], [""], [" \t\n"],
                         ["Mangal", "Mangal"], ["Noto *"], ["Noto ?"], ["Noto [Sans]"]):
            invalid.append({**self.policy, "approved_unicode_font_families": families})
        with self.policy_resource():
            for payload in invalid:
                with self.subTest(payload=payload):
                    self.path.write_text(json.dumps(payload), encoding="utf-8")
                    with self.assertRaisesRegex(ConfigurationError, "source-fonts policy"):
                        load_approved_unicode_font_families()

    def test_rejects_duplicates_after_normalization(self):
        self.policy["approved_unicode_font_families"] = ["Nirmala UI", " nirmala\tui "]
        self.write_policy()
        with self.policy_resource(), self.assertRaisesRegex(ConfigurationError, "duplicates a family"):
            load_approved_unicode_font_families()

    def test_missing_unreadable_malformed_and_duplicate_key_json_fail(self):
        with patch("gurubodh.legacy.font_detection.bundled_resource_path",
                   side_effect=BundledResourceError("missing resource")):
            with self.assertRaisesRegex(ConfigurationError, "Source-font policy.*missing"):
                load_approved_unicode_font_families()
        with self.policy_resource():
            self.path.unlink()
            with self.assertRaisesRegex(ConfigurationError, "Source-font policy.*missing"):
                load_approved_unicode_font_families()
            for data in (b"{", b"\xff", b'{"schema_version":"1.0.0","schema_version":"1.0.0",'
                         b'"approved_unicode_font_families":["Mangal"]}'):
                with self.subTest(data=data):
                    self.path.write_bytes(data)
                    with self.assertRaisesRegex(ConfigurationError, "Source-font policy.*malformed"):
                        load_approved_unicode_font_families()
            with patch.object(Path, "read_text", side_effect=PermissionError("unreadable")):
                with self.assertRaisesRegex(ConfigurationError, "Source-font policy.*unreadable"):
                    load_approved_unicode_font_families()

    def test_invalid_policy_never_reuses_previous_approvals_or_skips_fontless_sources(self):
        with self.policy_resource():
            validate_supported_source_fonts(self.source("Mangal"))
            self.path.write_text("{}")
            for family in ("Mangal", "APS-DV-Prakash", None):
                with self.subTest(family=family), self.assertRaises(ConfigurationError):
                    validate_supported_source_fonts(self.source(family))

    def test_missing_or_malformed_policy_schema_fails_closed(self):
        schema = self.root / "source-fonts.schema.json"
        for data in (None, "{", '{"type": "not-a-json-schema-type"}'):
            with self.subTest(data=data):
                if data is not None:
                    schema.write_text(data)
                _validator.cache_clear()
                self.addCleanup(_validator.cache_clear)
                with patch("gurubodh.schema_validation.schema_path", return_value=schema):
                    with self.assertRaisesRegex(ConfigurationError, "source-fonts policy.*Schema"):
                        load_approved_unicode_font_families()

    def test_cwd_and_project_environment_cannot_override_bundled_policy(self):
        fake = self.root / POLICY_PATH
        fake.parent.mkdir(parents=True)
        fake.write_text("{}")
        with chdir(self.root), patch.dict(os.environ, {"GURUBODH_CLI_ROOT": str(self.root)}):
            self.assertIn("mangal", load_approved_unicode_font_families())
            with self.assertRaises(UnsupportedSourceFontError):
                validate_supported_source_fonts(self.source("Unapproved Family"))

    def test_manifest_and_execution_profiles_reject_approval_overrides(self):
        for kind, relative in (
            ("subject-manifest", "tests/fixtures/job-components/subjects/unicode-bilingual.json"),
            ("proofreading-profile", "config/job-components/profiles/proofreading/gemini-3.6-flash-v1.json"),
            ("chunking-profile", "config/job-components/profiles/chunking/bge-m3-semantic-window-v1.json"),
        ):
            with self.subTest(kind=kind):
                payload = json.loads((CLI_ROOT / relative).read_text())
                payload["approved_unicode_font_families"] = ["Unapproved Family"]
                with self.assertRaisesRegex(ConfigurationError, "approved_unicode_font_families is not allowed"):
                    validate_component(payload, kind)

    def test_json_addition_is_accepted_by_both_commands_for_both_languages(self):
        from google.genai.models import Models

        for language in ("hi-IN", "mr-IN"):
            with self.subTest(language=language):
                root = self.root / language
                manifest, source = write_subject(root, CLI_ROOT / "tests/fixtures", "unicode", language)
                document = Document(source)
                document.paragraphs[1].runs[0].font.name = "New Unicode Family"
                document.save(source)
                selectors = ["--project-root", root / "project", "--subject", manifest["manifest_id"],
                             "--language", language, "--environment", "development", "--storage-profile", "local"]
                commands = (["prep-subject", *selectors],
                            ["lab", "proofread", "--source", source, "--locale", language,
                             "--lab-root", root / "lab", "--project-root", root / "project"])
                environment = {"GURUBODH_SOURCE_LIBRARY_ROOT": str(root / "source"),
                               "GURUBODH_CMS_LIBRARY_ROOT": str(root / "artifacts"),
                               "GEMINI_API_KEY": "offline-fixture-key"}
                with (self.policy_resource(), patch.dict(os.environ, environment),
                      patch.object(Models, "generate_content", side_effect=generate_content) as gemini,
                      patch("gurubodh.lab_proofread.convert_docx") as convert):
                    self.write_policy()
                    for command in commands:
                        if command[0] == "prep-subject":
                            invoke(command, error="Unicode-only source-font requirement")
                        else:
                            invoke(command, error="Unsupported source font family")
                    gemini.assert_not_called()
                    convert.assert_not_called()
                    self.assertFalse(list((root / "artifacts").glob("**/chapter_content_manifest.json")))
                    self.path.unlink()
                    for command in commands:
                        invoke(command, error="Source-font policy")
                    self.path.write_text("{}")
                    for command in commands:
                        invoke(command, error="source-fonts policy")
                    gemini.assert_not_called()
                    convert.assert_not_called()
                    self.assertFalse(list((root / "artifacts").glob("**/job-state.json")))
                    updated = {**self.policy, "approved_unicode_font_families": [
                        *self.policy["approved_unicode_font_families"], "new unicode family",
                    ]}
                    self.path.write_text(json.dumps(updated))
                    for command in commands:
                        invoke(command)
                    self.assertEqual(gemini.call_count, 2)
                    convert.assert_not_called()
                    self.assertEqual(len(list((root / "artifacts").glob("**/chapter_content_manifest.json"))), 1)
                    self.assertEqual(len(list((root / "lab").glob("**/succeeded/*/run_manifest.json"))), 1)
